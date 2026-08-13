"""Deterministic round-three coverage for the remaining data-facing services.

Every external boundary in this module is replaced with a local fake or mock.
The tests intentionally exercise low-frequency fallback, validation, and error
branches without live databases, Redis, Chroma, OpenAI, HTTP, filesystems, or
email services.
"""
from __future__ import annotations

import builtins
import json
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, Mock

import pandas as pd
import pytest

import app.services.auto_categorization_service as auto_mod
import app.services.daily_aggregation_service as aggregation_mod
import app.services.daily_summary_service as summary_mod
import app.services.email_service as email_mod
import app.services.entity_service as entity_mod
import app.services.excel_processing_service as processing_mod
import app.services.excel_validation_service as validation_mod
import app.services.exit_service as exit_mod
from app.models import UserType
from app.services.whatsapp_service import MessageResponse


class QueryFake:
    def __init__(self, values=(), first_value=None, count_value=None):
        self.values = list(values)
        self.first_value = first_value
        self.count_value = count_value

    def filter(self, *_args, **_kwargs):
        return self

    def between(self, *_args, **_kwargs):
        return self

    def all(self):
        return self.values

    def first(self):
        return self.first_value

    def count(self):
        return self.count_value if self.count_value is not None else len(self.values)


class ContextDB:
    def __init__(self, query=None, query_side_effect=None):
        self.query_value = query or QueryFake()
        self.query_side_effect = query_side_effect
        self.added = []
        self.commit = MagicMock()
        self.close = MagicMock()
        self.refresh = MagicMock()
        self.append_session_data = MagicMock()

    def query(self, *_args, **_kwargs):
        if self.query_side_effect:
            raise self.query_side_effect
        return self.query_value

    def add(self, value):
        self.added.append(value)

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False


class SequenceDB(ContextDB):
    def __init__(self, queries):
        super().__init__()
        self.queries = list(queries)

    def query(self, *_args, **_kwargs):
        return self.queries.pop(0)


class CollectionFake:
    def __init__(self, result=None, count=0):
        self.name = "category_items"
        self.result = result or {"documents": [[]], "metadatas": [[]], "distances": [[]]}
        self.count_value = count
        self.add = MagicMock()
        self.query_calls = 0
        self.query_side_effect = None

    def query(self, **_kwargs):
        self.query_calls += 1
        if self.query_side_effect:
            effect = self.query_side_effect
            if isinstance(effect, list):
                effect = effect.pop(0)
            if isinstance(effect, BaseException):
                raise effect
            return effect
        return self.result

    def count(self):
        if isinstance(self.count_value, BaseException):
            raise self.count_value
        return self.count_value


class ClientFake:
    def __init__(self, collection):
        self.collection = collection
        self.heartbeat = MagicMock(return_value=True)
        self.delete_collection = MagicMock()

    def get_or_create_collection(self, **_kwargs):
        return self.collection


def settings(**overrides):
    values = {
        "chroma_host": "localhost",
        "chroma_port": 8000,
        "enable_remote_categorization": False,
        "enable_daily_summarization": True,
        "email_templates_path": "templates",
        "support_email": "support@example.com",
        "email_signature": "Regards",
        "procucev_link": "https://procucev.example",
        "redis_session_storage_enabled": False,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def session(**overrides):
    values = {
        "session_id": "s1",
        "external_user_id": "+91 123-456",
        "phone_number": "+91123456",
        "user_type": UserType.buyer,
        "rfq_ids": [],
        "rfq_id": None,
        "product_items": [],
        "extracted_entities": {},
        "bfs_search_count": 0,
        "bfs_price_accepted": [],
        "products_bid_for": [],
        "bids_accepted": [],
        "rfqs_with_response": [],
        "seller_responses": [],
        "counter_offers_accepted": [],
        "counter_offers_made": [],
        "created_at": datetime(2025, 1, 1),
        "completed_at": None,
        "last_activity_at": datetime(2025, 1, 1, 0, 30),
        "session_state": None,
        "parent_session_id": None,
        "workflow_state": {},
        "conversation_history": {"messages": []},
        "workflow_type": "buy_something",
        "outcome": None,
        "retention_date": None,
        "extracted_entities": {},
    }
    values.update(overrides)
    return SimpleNamespace(**values)


# Auto categorization -------------------------------------------------------
def bare_auto(collection=None):
    service = auto_mod.AutoCategorizationService.__new__(auto_mod.AutoCategorizationService)
    service.collection = collection or CollectionFake()
    service.chroma_client = MagicMock()
    service.chroma_client.get_or_create_collection.return_value = service.collection
    service.embedding_function = "embedding"
    service.openai_service = SimpleNamespace()
    service.learning_service = SimpleNamespace()
    return service


def test_auto_categorization_fallback_population_query_and_health_edges(monkeypatch, tmp_path):
    module_file = tmp_path / "outside" / "auto.py"
    module_file.parent.mkdir()
    monkeypatch.setattr(auto_mod, "__file__", str(module_file))
    assert auto_mod.get_project_root() == module_file.parent.parent.parent

    client = ClientFake(CollectionFake())
    monkeypatch.setattr(auto_mod.embedding_functions, "SentenceTransformerEmbeddingFunction", lambda **_: "embed")
    monkeypatch.setattr(auto_mod.chromadb, "HttpClient", lambda **_: client)
    monkeypatch.setattr(auto_mod, "get_settings", lambda: settings())
    monkeypatch.setattr(auto_mod, "OpenAIService", lambda: object())
    monkeypatch.setattr(auto_mod, "LearningCategorizationService", lambda: object())
    client.heartbeat.side_effect = RuntimeError("chroma down")
    with pytest.raises(RuntimeError, match="ChromaDB server not available"):
        auto_mod.AutoCategorizationService()

    collection = CollectionFake()
    service = bare_auto(collection)
    service._get_local_category_data = MagicMock(return_value=[
        {"id": "1", "item": "pump", "category": "Tools", "division": None, "serial_no": None},
        {"id": "2", "item": "", "category": "Tools"},
        {"id": "3", "item": "bolt", "category": None},
    ])
    monkeypatch.setattr(auto_mod, "get_settings", lambda: settings())
    assert service.populate_embeddings_from_db() == 3
    payload = collection.add.call_args.kwargs
    assert payload["documents"] == ["pump"]
    assert "division" not in payload["metadatas"][0]
    assert "serial_no" not in payload["metadatas"][0]

    service._get_local_category_data.return_value = [{"id": "bad", "item": None, "category": None}]
    assert service.populate_embeddings_from_db() == 0

    failing_db = ContextDB(query_side_effect=RuntimeError("database"))
    monkeypatch.setattr(auto_mod, "get_db_session", lambda: failing_db)
    service._get_local_category_data = auto_mod.AutoCategorizationService._get_local_category_data.__get__(service)
    with pytest.raises(RuntimeError, match="database"):
        service._get_local_category_data()
    failing_db.close.assert_called_once()

    remote_settings = iter([settings(enable_remote_categorization=True), settings(enable_remote_categorization=False)])
    monkeypatch.setattr(auto_mod, "get_settings", lambda: next(remote_settings))
    service._get_remote_category_data = MagicMock(side_effect=RuntimeError("remote unavailable"))
    service._get_local_category_data = MagicMock(return_value=[])
    assert service.populate_embeddings_from_db() == 0

    empty = bare_auto(CollectionFake({"documents": [[]], "metadatas": [[]], "distances": [[]]}))
    assert empty._get_similar_items("none") == []
    empty.collection.query_side_effect = RuntimeError("unexpected query")
    with pytest.raises(Exception, match="Failed to get similar items"):
        empty._get_similar_items("broken")

    refreshed = CollectionFake({"documents": [[]], "metadatas": [[]], "distances": [[]]})
    stale = CollectionFake()
    stale.query_side_effect = [RuntimeError("404 collection does not exist")]
    empty.collection = stale
    empty._refresh_collection = MagicMock(side_effect=lambda: setattr(empty, "collection", refreshed))
    assert empty._query_collection("pump")["documents"] == [[]]
    empty.collection = CollectionFake()
    empty.collection.query_side_effect = RuntimeError("other failure")
    with pytest.raises(RuntimeError, match="other failure"):
        empty._query_collection("pump")

    empty.collection.count_value = 4
    monkeypatch.setattr(auto_mod, "get_settings", lambda: settings())
    assert empty.get_collection_stats()["total_items"] == 4
    empty.collection.count_value = RuntimeError("count failed")
    assert empty.get_collection_stats()["error"] == "count failed"

    empty.get_collection_stats = MagicMock(return_value={"total_items": 2})
    healthy_db = ContextDB(QueryFake(count_value=5))
    monkeypatch.setattr(auto_mod, "get_db_session", lambda: healthy_db)
    empty.openai_service = SimpleNamespace(client=object())
    assert empty.health_check()["status"] == "healthy"
    unhealthy_db = ContextDB(query_side_effect=RuntimeError("db down"))
    monkeypatch.setattr(auto_mod, "get_db_session", lambda: unhealthy_db)
    assert empty.health_check()["status"] == "unhealthy"
    empty.get_collection_stats = MagicMock(side_effect=RuntimeError("stats down"))
    assert empty.health_check()["status"] == "unhealthy"
    empty.get_collection_stats = MagicMock(return_value={"total_items": 1})
    empty.openai_service = SimpleNamespace()
    monkeypatch.setattr(auto_mod, "get_db_session", lambda: healthy_db)
    assert empty.health_check()["status"] == "unhealthy"


# Daily aggregation and summary --------------------------------------------
def test_daily_aggregation_optional_shapes_rolling_and_storage(monkeypatch):
    service = aggregation_mod.DailyAggregationService.__new__(aggregation_mod.DailyAggregationService)
    buyer_list = session(
        external_user_id="b1", rfq_id="rfq-one", product_items=[{"category": "Tools"}],
        products_bid_for=["p1", "p2"], bids_accepted=["b1"], rfqs_with_response=["r", "r"],
        bfs_search_count=2, products_searched_count=3, total_rfq_responses_received=4,
    )
    buyer_dict = session(
        external_user_id="b2", rfq_ids=["rfq-two", None], product_items=[{"category": "Hardware"}],
        products_bid_for={"products": ["p3"]}, bids_accepted={"count": 2},
        rfqs_with_response={"unique_rfqs": 3}, bfs_search_count=1,
    )
    buyer_empty = session(external_user_id="b3", product_items={}, products_bid_for={}, bids_accepted={}, rfqs_with_response={})
    metrics = service._calculate_buyer_summary_metrics([buyer_list, buyer_dict, buyer_empty], date.today())
    assert metrics["total_rfqs_submitted"] == 2
    assert metrics["products_bid_for"] == 3
    assert metrics["bids_accepted_by_buyers"] == 3
    assert metrics["rfqs_with_at_least_one_response"] == 4
    assert metrics["bfs_products_searched"] == 3

    seller_list = session(
        user_type=UserType.seller, external_user_id="s1", rfq_id="rfq-s1",
        seller_responses=["r1"], bids_accepted=["b1"], counter_offers_accepted=["c1", "c2"],
        counter_offers_made=["m1"], extracted_entities=[{"category": "Tools"}],
    )
    seller_dict = session(
        user_type=UserType.seller, external_user_id="s2", rfq_ids=["rfq-s2"],
        seller_responses={"count": 2}, bids_accepted={"count": 3},
        counter_offers_accepted={"count": 4}, counter_offers_made={"count": 5},
        extracted_entities={"category": "Hardware"},
    )
    seller_empty = session(user_type=UserType.seller, seller_responses={}, bids_accepted={}, counter_offers_accepted={})
    seller_metrics = service._calculate_seller_summary_metrics([seller_list, seller_dict, seller_empty], date.today())
    assert seller_metrics["rfqs_requested"] == 2
    assert seller_metrics["counter_offers_accepted"] == 6

    category_buyer = session(
        user_type=UserType.buyer, rfq_id="rfq-category", product_items=[{"category": "Tools"}],
        rfqs_with_response=["r1"], bfs_search_count=2, bfs_price_accepted=["p"],
        extracted_entities={"category": "Tools"},
    )
    category_seller = session(
        user_type=UserType.seller, product_items=[{"category": "Hardware"}],
        rfqs_with_response={"count": 2}, bfs_price_accepted={"count": 3},
        counter_offers_accepted={"count": 4}, counter_offers_made={"count": 5},
        extracted_entities=[{"description": "Hardware"}],
    )
    categories = service._calculate_category_summary_metrics([category_buyer, category_seller], date.today())
    assert categories["Tools"]["rfqs_uploaded"] == 1
    assert categories["Hardware"]["counter_offers_accepted_by_sellers"] == 4
    assert categories["Hardware"]["new_counter_offer_by_seller"] == 5
    assert service._extract_categories_from_session(session(product_items={}, extracted_entities={})) == []

    db = ContextDB(QueryFake(first_value=None))
    service._store_metrics(db, date.today(), "buyer_summary", {"x": 1})
    assert len(db.added) == 1
    existing = SimpleNamespace(metric_data={}, is_complete=False)
    service._store_metrics(ContextDB(QueryFake(first_value=existing)), date.today(), "buyer_summary", {"x": 2})
    assert existing.metric_data == {"x": 2} and existing.is_complete

    metric_rows = [
        SimpleNamespace(metric_type="buyer_summary", metric_data={"total_rfqs_submitted": 2, "avg_products_per_rfq": 2}),
        SimpleNamespace(metric_type="seller_summary", metric_data={"seller_chats_initiated": 1}),
        SimpleNamespace(metric_type="category_summary", metric_data={"Tools": {"rfqs_uploaded": 1}}),
        SimpleNamespace(metric_type="other", metric_data={}),
    ]
    roll_db = ContextDB(QueryFake(values=metric_rows, first_value=None))
    monkeypatch.setattr(aggregation_mod, "get_db_session", lambda: roll_db)
    service._calculate_rolling_window(date(2025, 1, 1), "7day", 7)
    assert len(roll_db.added) == 1
    existing_window = SimpleNamespace()
    roll_db_existing = ContextDB(QueryFake(values=metric_rows, first_value=existing_window))
    monkeypatch.setattr(aggregation_mod, "get_db_session", lambda: roll_db_existing)
    service._calculate_rolling_window(date(2025, 1, 1), "7day", 7)
    assert existing_window.buyer_metrics["total_rfqs_submitted"] == 2
    service._calculate_rolling_window = MagicMock(side_effect=RuntimeError("rolling"))
    with pytest.raises(RuntimeError):
        service._update_rolling_windows(date.today())


@pytest.mark.asyncio
async def test_daily_summary_sparse_optional_fields_and_lookup_edges(monkeypatch):
    service = summary_mod.DailySummaryService.__new__(summary_mod.DailySummaryService)
    service.settings = settings(enable_daily_summarization=False)
    assert await service.generate_daily_summary("u") is None
    service.settings = settings(enable_daily_summarization=True)

    sparse = session(
        user_type=None, rfq_ids=[], rfq_id=None, product_items=[], seller_responses=[],
        created_at=None, completed_at=None, last_activity_at=None, session_state=None,
        parent_session_id=None, extracted_entities=[{"description": "pump"}, {"category": "Tools"}],
    )
    summary_dict = SimpleNamespace(extracted_entities={"description": "bolt", "category": "Hardware"})
    summary_list = SimpleNamespace(extracted_entities=[{"description": "nut"}, {"category": "Fasteners"}])
    db = SequenceDB([QueryFake(values=[sparse]), QueryFake(values=[summary_dict, summary_list]), QueryFake(first_value=None)])
    monkeypatch.setattr(summary_mod, "get_db_session", lambda: db)
    created = await service.generate_daily_summary("u", date(2025, 1, 1))
    assert created is not None
    assert created.avg_session_duration is None
    assert set(created.primary_product_categories) == {"pump", "Tools", "bolt", "Hardware", "nut", "Fasteners"}

    existing = SimpleNamespace()
    full = session(
        rfq_ids=["r1"], product_items=[{"x": 1}], seller_responses=["response"],
        session_state=SimpleNamespace(value="completed"), parent_session_id="parent",
        extracted_entities={"description": "pump", "category": "Tools"},
        completed_at=datetime(2025, 1, 1, 1), last_activity_at=datetime(2025, 1, 1, 1),
    )
    db_existing = SequenceDB([QueryFake(values=[full]), QueryFake(values=[]), QueryFake(first_value=existing)])
    monkeypatch.setattr(summary_mod, "get_db_session", lambda: db_existing)
    assert await service.generate_daily_summary("u", date(2025, 1, 1)) is existing
    assert existing.sessions_count == 1

    db_none = ContextDB(QueryFake(values=[]))
    monkeypatch.setattr(summary_mod, "get_db_session", lambda: db_none)
    assert await service.get_user_daily_summary("u", date(2025, 1, 1)) is None
    optional = SimpleNamespace(
        date=date(2025, 1, 1), user_type=None, sessions_count=0, rfqs_created_count=0,
        session_states_breakdown=None, seller_interaction_count=0, avg_session_duration=None,
        primary_product_categories=None, product_items_count=0, unique_categories_count=0,
        session_continuation_count=0, total_rfq_value_estimate=0,
    )
    monkeypatch.setattr(summary_mod, "get_db_session", lambda: ContextDB(QueryFake(first_value=optional)))
    result = await service.get_user_daily_summary("u", date(2025, 1, 1))
    assert result["user_type"] is None and result["total_rfq_value_estimate"] is None
    assert service._calculate_duration(session(created_at="bad", completed_at="bad", last_activity_at=None)) is None


# Email ---------------------------------------------------------------------
class FileContext:
    def __enter__(self):
        return self

    def __exit__(self, *_args):
        return False


def email_service():
    service = email_mod.EmailService.__new__(email_mod.EmailService)
    service.settings = settings()
    service.templates_cache = {}
    service.email_api = SimpleNamespace(send_email=AsyncMock())
    return service


def test_email_template_file_role_substitution_and_listing_edges(monkeypatch):
    service = email_service()
    exists = Mock(side_effect=[True, False, False, False, False])
    monkeypatch.setattr(email_mod.os.path, "exists", exists)
    monkeypatch.setattr(builtins, "open", Mock(side_effect=FileNotFoundError("race")))
    assert service._load_template("missing") is None

    invalid_file = FileContext()
    monkeypatch.setattr(email_mod.os.path, "exists", Mock(side_effect=[True, False, False, False, False]))
    monkeypatch.setattr(builtins, "open", Mock(return_value=invalid_file))
    monkeypatch.setattr(email_mod.json, "load", Mock(side_effect=json.JSONDecodeError("bad", "", 0)))
    assert service._load_template("invalid") is None

    monkeypatch.setattr(email_mod.os.path, "exists", Mock(side_effect=[True, False, False, False, False]))
    monkeypatch.setattr(builtins, "open", Mock(side_effect=RuntimeError("read failure")))
    assert service._load_template("broken") is None

    buyer = service._process_template({"supports_buyer": False, "to": "a@example.com", "subject": "Hi", "body": "Body"}, {}, "buyer")
    seller = service._process_template({"supports_seller": False, "to": "a@example.com", "subject": "Hi", "body": "Body"}, {}, "seller")
    assert buyer["to"] == seller["to"] == ["a@example.com"]
    assert service._substitute_variables("unmatched {", {}) == "unmatched {"

    service._load_template = Mock(side_effect=[
        {"scenario_id": "A", "subject": "A", "supports_buyer": False},
        None,
    ])
    monkeypatch.setattr(email_mod.os, "listdir", Mock(return_value=["a.json", "notes.txt", "b.json"]))
    assert service.list_available_templates() == [{"name": "a", "scenario_id": "A", "subject": "A", "supports_buyer": False, "supports_seller": True}]
    monkeypatch.setattr(email_mod.os, "listdir", Mock(side_effect=OSError("directory unavailable")))
    assert service.list_available_templates() == []


# Entity extraction ---------------------------------------------------------
def entity_service(ai=None):
    ai = ai or SimpleNamespace(
        extract_entities=AsyncMock(),
        validate_delivery_date=AsyncMock(),
        extract_entities_with_summary_context=AsyncMock(),
        merge_resolved_references_with_entities=AsyncMock(),
    )
    return entity_mod.EntityService(ai), ai


@pytest.mark.asyncio
async def test_entity_modification_state_shapes_global_fields_and_operations():
    service, ai = entity_service()
    ai.extract_entities.return_value = {
        "is_modification_extraction": True,
        "has_new_values": False,
        "modifications": [],
        "modification_intent": "delivery address",
        "confidence": 60,
    }
    products = [{"entities": {"description": "pump", "quantity": 1}}]
    states = [
        {"pending_combined_rfq": {"products": products}},
        {"pending_rfq": products[0]},
        {"pending_optional_combined_rfq": {"products": products}},
        {"pending_optional_rfq": products[0]},
        {"incomplete_products": products},
        {"complete_products": products},
        {"extracted_entities": [{"description": "pump", "attachments": ["file"]}]},
    ]
    for workflow_state in states:
        result = await service._handle_modification_extraction("change address", {"workflow_state": workflow_state}, "modification_request")
        assert result["requires_clarification"]

    service._handle_standard_extraction = AsyncMock(return_value={"products": [{"description": "fallback"}]})
    assert (await service._handle_modification_extraction("nothing", {"workflow_state": {}}, "modification_request"))["products"][0]["description"] == "fallback"

    future = (datetime.now() + timedelta(days=5)).strftime("%Y-%m-%d")
    ai.extract_entities.return_value = {
        "is_modification_extraction": True,
        "has_new_values": True,
        "modifications": [{"operation_type": "modify", "target_product_index": 0, "new_quantity": 2}],
        "deliveryDate": future,
        "state": "MH",
        "city": "Pune",
        "pincode": "411005",
        "confidence": 90,
    }
    service._apply_modifications_to_existing_products = MagicMock(return_value=[{"description": "pump", "quantity": 2}])
    service._validate_dates_in_products = AsyncMock(return_value=([{"description": "pump", "quantity": 2}], False))
    service._clean_invalid_descriptions = MagicMock(side_effect=lambda value: value)
    service._auto_fill_location_from_pincode = AsyncMock(side_effect=lambda value: value)
    result = await service._handle_modification_extraction("move it", {"workflow_state": {"pending_rfq": products[0]}}, "modification_request")
    assert result["success"] and result["is_modification"]
    assert service._apply_modifications_to_existing_products.call_args.args[3] == {
        "deliveryDate": future, "state": "MH", "city": "Pune", "pincode": "411005"
    }

    ai.extract_entities.return_value = {
        "is_modification_extraction": True,
        "has_new_values": False,
        "modifications": [{"operation_type": "remove", "target_product_index": 0}],
        "confidence": 80,
    }
    service._apply_modifications_to_existing_products.return_value = []
    service._validate_dates_in_products.return_value = ([], False)
    removed = await service._handle_modification_extraction("remove it", {"workflow_state": {"pending_rfq": products[0]}}, "modification_request")
    assert removed["success"] and removed["products"] == []

    existing = [
        {"entities": {"description": "pump", "city": "Pune", "state": "MH", "date_validation_error": "old"}},
        {"entities": {"description": "bolt", "city": "Pune", "state": "MH"}},
    ]
    service._apply_modifications_to_existing_products = entity_mod.EntityService._apply_modifications_to_existing_products.__get__(service)
    applied = service._apply_modifications_to_existing_products(existing, [
        {"operation_type": "modify", "target_product_index": None, "new_quantity": 1},
        {"operation_type": "modify", "target_product_index": 99, "new_quantity": 1},
        {"operation_type": "remove", "target_product_index": None},
        {"operation_type": "remove", "target_product_index": 99},
        {"operation_type": "add"},
        {"operation_type": "unknown", "target_product_index": 0},
        {"operation_type": "modify", "target_product_index": 0, "delivery_date": future},
        {"operation_type": "add", "target_product_description": "nut", "new_description": "nut", "new_quantity": 3},
        {"operation_type": "remove", "target_product_index": 1},
    ], "change", {"city": "Mumbai"})
    assert len(applied) == 3 and applied[0]["city"] == "Mumbai" and applied[2]["description"] == "nut"
    only_existing = [{"entities": {"description": "plain"}}]
    assert len(service._apply_modifications_to_existing_products(only_existing, [{"operation_type": "add"}], "change")) == 1


@pytest.mark.asyncio
async def test_entity_date_summary_merge_and_pincode_edges(monkeypatch):
    service, ai = entity_service()
    future = (datetime.now() + timedelta(days=3)).strftime("%Y-%m-%d")
    ai.validate_delivery_date.side_effect = [
        {"is_valid": True, "normalized_date": future, "validation_issues": ["warning"]},
        {"is_valid": True, "normalized_date": future},
        {"is_valid": True, "normalized_date": "2000-01-01"},
        {"is_valid": True, "normalized_date": "not-a-date"},
        {"is_valid": True, "normalized_date": future, "validation_issues": ["warning"]},
        {"is_valid": True, "normalized_date": "2000-01-01"},
        {"is_valid": True, "normalized_date": "not-a-date"},
        {"is_valid": True, "normalized_date": future},
    ]
    invalid, has_error = await service._validate_dates_in_products(
        [{"deliveryDate": future, "date_validation_error": "old"}], "date"
    )
    assert has_error and invalid[0]["deliveryDate"] is None
    valid, has_error = await service._validate_dates_in_products([{"deliveryDate": future, "date_validation_error": "old"}], "date")
    assert not has_error and "date_validation_error" not in valid[0]
    past, has_error = await service._validate_dates_in_products([{"deliveryDate": future}], "date")
    assert has_error and past[0]["deliveryDate"] is None
    malformed, has_error = await service._validate_dates_in_products([{"deliveryDate": future}], "date")
    assert has_error and malformed[0]["deliveryDate"] is None

    for expected_error, result in [(True, {"deliveryDate": future, "description": "x"}), (True, {"deliveryDate": future, "description": "x"}), (True, {"deliveryDate": future, "description": "x"}), (False, {"deliveryDate": future})]:
        validated, error = await service._validate_date_in_entity(result, "date")
        assert error is expected_error
        if error:
            assert validated["deliveryDate"] is None
    no_date, no_error = await service._validate_date_in_entity({"description": "x"}, "date")
    assert no_date == {"description": "x"} and not no_error

    service._handle_reference_extraction = AsyncMock(return_value={"reference": True})
    reference_context = {"intent_result": {"intent": "reference_request", "confidence": 90, "reasoning": "history", "context_analysis": {"reference_details": {"reference_type": "address", "has_history": True}}}}
    assert await service.extract_entities_with_summary_context("usual", reference_context) == {"reference": True}
    service._handle_modification_extraction = AsyncMock(return_value={"modification": True})
    assert await service.extract_entities_with_summary_context("change", {"chat_summaries": [1]}, "modification_request") == {"modification": True}
    service._handle_standard_extraction = AsyncMock(return_value={"standard": True})
    assert await service.extract_entities_with_summary_context("new", {}) == {"standard": True}

    ai.extract_entities_with_summary_context.return_value = {
        "products": [{"description": "pump"}],
        "resolved_references": [{"reference_phrase": "usual", "resolved_field": "city", "resolved_value": "Pune"}],
    }
    ai.merge_resolved_references_with_entities.return_value = {"success": True, "updated_products": [{"description": "pump", "city": "Pune"}]}
    result = await service.extract_entities_with_summary_context("usual pump", {"chat_summaries": [{"text": "old"}]}, "buy_something")
    assert result["products"][0]["city"] == "Pune"
    ai.extract_entities_with_summary_context.return_value = {"success": True}
    assert await service.extract_entities_with_summary_context("empty", {"chat_summaries": [1]}) == {"success": True}
    ai.extract_entities_with_summary_context.side_effect = RuntimeError("summary down")
    ai.extract_entities.side_effect = None
    ai.extract_entities.return_value = {"products": []}
    service._handle_standard_extraction = AsyncMock(return_value={"products": []})
    fallback = await service.extract_entities_with_summary_context("fallback", {"chat_summaries": [1]})
    assert fallback == {"products": []}

    base = [{"description": None, "unitofMeasures": "kg"}, {"description": "bolt", "quantity": 1}]
    positional = service._merge_new_extraction_with_existing_products(
        base, [{"description": "pump", "unitofMeasures": "unit(s)", "quantity": 2}, {"description": "nut", "quantity": 3}], "more"
    )
    assert positional[0]["description"] == "pump" and positional[0]["unitofMeasures"] == "kg"
    merged = service._merge_new_extraction_with_existing_products(
        [{"description": "bolt", "unitofMeasures": "kg"}],
        [{"description": "bolt", "quantity": 2, "unitofMeasures": "unit(s)"}, {"description": "NO_PRODUCTS_MENTIONED"}, {"description": None, "city": "Pune"}, {"description": 4, "state": "MH"}, {"description": ""}, {"description": "nut"}],
        "more",
    )
    assert merged[0]["quantity"] == 2 and merged[0]["city"] == "Pune" and len(merged) == 2

    assert await service._auto_fill_location_from_pincode([]) == []
    assert await service._auto_fill_location_from_pincode([{"description": "x"}]) == [{"description": "x"}]
    monkeypatch.setattr(entity_mod, "get_location_from_pincode_async", AsyncMock(return_value={"city": "Pune", "state": "MH"}))
    located = await service._auto_fill_location_from_pincode([{"pincode": "411005", "city": "old"}, {"pincode": "411005"}])
    assert all(item["city"] == "Pune" for item in located)
    monkeypatch.setattr(entity_mod, "get_location_from_pincode_async", AsyncMock(side_effect=RuntimeError("lookup")))
    failed = await service._auto_fill_location_from_pincode([{"pincode": "411005", "city": "old"}])
    assert failed[0]["city"] == "old"
    monkeypatch.setattr(entity_mod, "get_location_from_pincode_async", AsyncMock(return_value=None))
    missing = await service._auto_fill_location_from_pincode([{"pincode": "411005"}])
    assert missing[0]["pincode"] is None


# Excel processing ----------------------------------------------------------
def processing_service():
    return processing_mod.ExcelProcessingService(SimpleNamespace(process_excel_to_rfqs=AsyncMock()))


def configure_processing(monkeypatch, service, frame=None):
    monkeypatch.setattr(service, "_is_valid_excel_file", lambda *_: True)
    service._validate_excel_structure = AsyncMock(return_value={"valid": True})
    monkeypatch.setattr(processing_mod.pd, "ExcelFile", lambda *_: SimpleNamespace(sheet_names=["Sheet1"], close=MagicMock()))
    monkeypatch.setattr(processing_mod.pd, "read_excel", lambda *_args, **_kwargs: frame if frame is not None else pd.DataFrame([["pump", "steel", 2]]))


@pytest.mark.asyncio
async def test_excel_processing_ingestion_openai_mapping_preprocess_and_structure_edges(monkeypatch):
    service = processing_service()
    assert (await service.process_excel_file(b"x" * (3 * 1024 * 1024 + 1), "items.xlsx"))["success"] is False
    monkeypatch.setattr(service, "_is_valid_excel_file", lambda *_: False)
    assert (await service.process_excel_file(b"data", "items.xlsx"))["success"] is False
    configure_processing(monkeypatch, service, pd.DataFrame([["pump", "steel", 2]]))
    service._validate_excel_structure.return_value = {"valid": False, "error": "rows"}
    assert (await service.process_excel_file(b"data", "items.xlsx"))["error"] == "rows"
    configure_processing(monkeypatch, service, pd.DataFrame([["pump", "steel", 2]]))
    monkeypatch.setattr(processing_mod.pd, "ExcelFile", lambda *_: SimpleNamespace(sheet_names=["A", "B"], close=MagicMock()))
    assert (await service.process_excel_file(b"data", "items.xlsx"))["success"] is False
    configure_processing(monkeypatch, service, pd.DataFrame([["pump", "steel", 2]]))
    monkeypatch.setattr(processing_mod.pd, "read_excel", Mock(side_effect=RuntimeError("read")))
    assert "Failed to read Excel" in (await service.process_excel_file(b"data", "items.xlsx"))["error"]
    configure_processing(monkeypatch, service, pd.DataFrame())
    assert (await service.process_excel_file(b"data", "items.xlsx"))["success"] is False
    configure_processing(monkeypatch, service, pd.DataFrame([["pump", "steel", 2]]))
    service._process_excel_with_openai = AsyncMock(return_value={"success": False, "error": "ai"})
    assert (await service.process_excel_file(b"data", "items.xlsx"))["error"] == "ai"
    configure_processing(monkeypatch, service, pd.DataFrame([["pump", "steel", 2]]))
    service._process_excel_with_openai = AsyncMock(return_value={
        "success": True,
        "rfqs": [{"products": [{"description": "pump", "brand": "steel", "quantity": 2}], "deliveryDate": "2025-01-01", "city": "Pune"}],
        "processing_summary": {"total_products": 0, "total_products_extracted": 1, "skipped_rows": 1, "skipped_items_summary": "one row"},
    })
    successful = await service.process_excel_file(b"data", "items.xlsx")
    assert successful["success"] and "rows skipped" in successful["summary"]
    configure_processing(monkeypatch, service, pd.DataFrame([["pump", "steel", 2]]))
    service._process_excel_with_openai = AsyncMock(return_value={
        "success": True,
        "rfqs": [{"products": [{"description": "pump", "quantity": 2}], "city": "Pune"}, {"products": [], "city": "Mumbai"}],
        "processing_summary": {},
    })
    assert (await service.process_excel_file(b"data", "items.xlsx"))["success"] is False

    service._process_excel_with_openai = processing_mod.ExcelProcessingService._process_excel_with_openai.__get__(service)
    service.openai_service.process_excel_to_rfqs = AsyncMock(return_value={"success": True, "rfqs": []})
    assert (await service._process_excel_with_openai(pd.DataFrame({"A": [1]}), "items.xlsx"))["success"]
    service.openai_service.process_excel_to_rfqs.return_value = {"success": False, "error": "bad"}
    assert not (await service._process_excel_with_openai(pd.DataFrame({"A": [1]}), "items.xlsx"))["success"]
    service.openai_service.process_excel_to_rfqs.side_effect = RuntimeError("openai")
    assert not (await service._process_excel_with_openai(pd.DataFrame({"A": [1]}), "items.xlsx"))["success"]

    assert service._create_fallback_mapping(["ItemDescription", "not-a-column"]) == {"ItemDescription": "ItemDescription"}
    df = pd.DataFrame([[None, None, None], ["pump", "steel", 2], ["bolt", None, 3]], columns=["ItemDescription", "Specification", "Quantity"])
    extracted = service._extract_items_with_mapping(df, list(df.columns), {c: c for c in df.columns})
    assert extracted["removed_rows"] == 1 and extracted["items"][0]["Uom"] == "unit(s)"
    class BadRows:
        def iterrows(self):
            raise RuntimeError("rows")
    assert not service._extract_items_with_mapping(BadRows(), [], {})["success"]
    assert service._validate_items_comprehensive([{"ItemDescription": "", "Quantity": "Five,", "Uom": "each"}])["valid"] is False
    clean = service._preprocess_excel_data(pd.DataFrame([[1, 1, None], [1, 1, None]], columns=["A", "A.1", "Unnamed: 2"]))
    assert list(clean.columns) == ["A"] and len(clean) == 1
    class BadFrame:
        @property
        def columns(self):
            raise RuntimeError("columns")
    bad = BadFrame()
    assert service._preprocess_excel_data(bad) is bad
    assert service._format_date_for_display("") == ""
    assert service._format_date_for_display("2025-01-02T00:00:00") == "02 January"
    assert service._format_date_for_display("2025-01-02") == "02 January"
    assert service._format_date_for_display("not-a-date") == "not-a-date"
    inconsistent = service._validate_rfqs_consistency([{"deliveryDate": "2025-01-01", "city": "A"}, {"deliveryDate": "2025-01-02", "city": "B"}, {"deliveryDate": "2025-01-03", "city": "C"}])
    assert inconsistent["valid"] is False and "+1 more" in inconsistent["error"]

    service._validate_excel_structure = processing_mod.ExcelProcessingService._validate_excel_structure.__get__(service)
    import openpyxl
    workbook = SimpleNamespace(active=SimpleNamespace(iter_rows=lambda: [[SimpleNamespace(value="x")]], merged_cells=SimpleNamespace(ranges=[])), close=MagicMock())
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_args, **_kwargs: workbook)
    assert (await service._validate_excel_structure(b"data"))["valid"]
    row_book = SimpleNamespace(active=SimpleNamespace(iter_rows=lambda: [[SimpleNamespace(value="x")] for _ in range(51)], merged_cells=SimpleNamespace(ranges=[])), close=MagicMock())
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *_args, **_kwargs: row_book)
    assert (await service._validate_excel_structure(b"data"))["valid"] is False
    monkeypatch.setattr(openpyxl, "load_workbook", Mock(side_effect=RuntimeError("openpyxl")))
    monkeypatch.setattr(processing_mod.pd, "read_excel", lambda *_args, **_kwargs: pd.DataFrame({"A": range(51)}))
    assert (await service._validate_excel_structure(b"data"))["valid"] is False
    monkeypatch.setattr(processing_mod.pd, "read_excel", Mock(side_effect=RuntimeError("pandas")))
    assert (await service._validate_excel_structure(b"data"))["valid"] is False

    output_error = processing_mod.io.BytesIO
    monkeypatch.setattr(processing_mod.io, "BytesIO", Mock(side_effect=RuntimeError("output")))
    with pytest.raises(RuntimeError, match="output"):
        service.create_standard_template([{"S.No": "bad", "Quantity": "bad"}])
    monkeypatch.setattr(processing_mod.io, "BytesIO", output_error)
    assert service.create_standard_template([{"S.No": "bad", "Quantity": "bad"}]).startswith(b"PK")


# Excel validation ----------------------------------------------------------
@pytest.mark.asyncio
async def test_excel_validation_pipeline_readers_structure_quality_and_type_edges(monkeypatch):
    service = validation_mod.ExcelValidationService()
    service._download_file_with_retry = AsyncMock(return_value=b"PK\x03\x04payload")
    service._validate_excel_readability = AsyncMock(return_value={"valid": True})
    service._validate_excel_structure_comprehensive = AsyncMock(return_value={"valid": True})
    service._validate_data_quality = AsyncMock(return_value={"valid": True})
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["valid"]
    service._download_file_with_retry.return_value = b"x" * (service.MAX_FILE_SIZE + 1)
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "file_too_large"
    service._download_file_with_retry.return_value = b"short"
    service._validate_file_integrity = Mock(return_value={"valid": False, "error_type": "corrupted_file"})
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "corrupted_file"
    service._validate_file_integrity = validation_mod.ExcelValidationService._validate_file_integrity.__get__(service)
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "corrupted_file"
    service._download_file_with_retry.return_value = b"not-excel"
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "invalid_format"
    service._download_file_with_retry.return_value = b"PK\x03\x04payload"
    service._validate_excel_readability.return_value = {"valid": False, "error_type": "unreadable_file"}
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "unreadable_file"
    service._validate_excel_readability.return_value = {"valid": True}
    service._validate_excel_structure_comprehensive.return_value = {"valid": False, "error_type": "structure_validation_error"}
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "structure_validation_error"
    service._validate_excel_structure_comprehensive.return_value = {"valid": True}
    service._validate_data_quality.return_value = {"valid": False, "error_type": "data_quality_error"}
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "data_quality_error"
    service._validate_file_integrity = Mock(side_effect=RuntimeError("outer"))
    assert (await service.validate_excel_file_from_url("url", "x.xlsx"))["error_type"] == "validation_error"

    service._validate_file_integrity = validation_mod.ExcelValidationService._validate_file_integrity.__get__(service)
    service._validate_excel_readability = validation_mod.ExcelValidationService._validate_excel_readability.__get__(service)
    service._validate_excel_structure_comprehensive = validation_mod.ExcelValidationService._validate_excel_structure_comprehensive.__get__(service)
    service._validate_data_quality = validation_mod.ExcelValidationService._validate_data_quality.__get__(service)
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(side_effect=ValueError("encrypted workbook")))
    assert (await service._validate_excel_readability(b"data", "x.xlsx"))["error_type"] == "password_protected"
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(side_effect=ValueError("bad engine")))
    workbook = SimpleNamespace(close=MagicMock())
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=workbook))
    assert (await service._validate_excel_readability(b"data", "x.xlsx"))["valid"]
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(side_effect=validation_mod.InvalidFileException("encrypted")))
    assert (await service._validate_excel_readability(b"data", "x.xlsx"))["error_type"] == "password_protected"
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(side_effect=validation_mod.InvalidFileException("bad")))
    monkeypatch.setattr(validation_mod.xlrd, "open_workbook", Mock(return_value=object()))
    assert (await service._validate_excel_readability(b"data", "x.xls"))["valid"]
    monkeypatch.setattr(validation_mod.xlrd, "open_workbook", Mock(side_effect=RuntimeError("bad")))
    assert (await service._validate_excel_readability(b"data", "x.xls"))["error_type"] == "unreadable_file"
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(side_effect=RuntimeError("outer")))
    assert (await service._validate_excel_readability(b"data", "x.xlsx"))["error_type"] == "readability_error"

    worksheet = SimpleNamespace(iter_rows=lambda: [[SimpleNamespace(value="h")]], merged_cells=SimpleNamespace(ranges=[]), row_dimensions={}, column_dimensions={}, _pivots=[])
    book = SimpleNamespace(worksheets=[worksheet], active=worksheet, close=MagicMock())
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=book))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["valid"]
    hidden = SimpleNamespace(iter_rows=lambda: [[SimpleNamespace(value="h")]], merged_cells=SimpleNamespace(ranges=[]), row_dimensions={1: SimpleNamespace(hidden=True)}, column_dimensions={"A": SimpleNamespace(hidden=True)}, _pivots=[])
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=SimpleNamespace(worksheets=[hidden], active=hidden, close=MagicMock())))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["valid"]
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=SimpleNamespace(worksheets=[worksheet, worksheet], active=worksheet, close=MagicMock())))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["error_type"] == "multiple_worksheets"
    many = SimpleNamespace(iter_rows=lambda: [[SimpleNamespace(value="h")] for _ in range(51)], merged_cells=SimpleNamespace(ranges=[]), row_dimensions={}, column_dimensions={}, _pivots=[])
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=SimpleNamespace(worksheets=[many], active=many, close=MagicMock())))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["error_type"] == "too_many_rows"
    merged = SimpleNamespace(iter_rows=lambda: [[SimpleNamespace(value="h")]], merged_cells=SimpleNamespace(ranges=["A1:B1"]), row_dimensions={}, column_dimensions={}, _pivots=[])
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=SimpleNamespace(worksheets=[merged], active=merged, close=MagicMock())))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["error_type"] == "merged_cells_found"
    pivot = SimpleNamespace(iter_rows=lambda: [[SimpleNamespace(value="h")]], merged_cells=SimpleNamespace(ranges=[]), row_dimensions={}, column_dimensions={}, _pivots=[object()])
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=SimpleNamespace(worksheets=[pivot], active=pivot, close=MagicMock())))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["error_type"] == "pivot_tables_found"
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(side_effect=RuntimeError("openpyxl")))
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=pd.DataFrame({"A": range(51)})))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["error_type"] == "too_many_rows"
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(side_effect=RuntimeError("pandas")))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["error_type"] == "structure_validation_error"

    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=pd.DataFrame()))
    assert (await service._validate_data_quality(b"data"))["error_type"] == "no_data_found"
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=pd.DataFrame({"A": [1], "B": [2]})))
    assert (await service._validate_data_quality(b"data"))["error_type"] == "no_headers_detected"
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=pd.DataFrame({"Item@": ["x"], "Qty": ["bad"]})))
    assert (await service._validate_data_quality(b"data"))["error_type"] == "invalid_header_characters"
    original_pattern = validation_mod.ExcelValidationService.SPECIAL_CHARS_PATTERN
    monkeypatch.setattr(validation_mod.ExcelValidationService, "SPECIAL_CHARS_PATTERN", r"$^")
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=pd.DataFrame({"数量": ["x"], "項目": ["y"]})))
    assert (await service._validate_data_quality(b"data"))["error_type"] == "non_english_headers"
    monkeypatch.setattr(validation_mod.ExcelValidationService, "SPECIAL_CHARS_PATTERN", original_pattern)
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=pd.DataFrame({"ItemDescription": ["x"], "Quantity": ["bad"]})))
    assert (await service._validate_data_quality(b"data"))["error_type"] == "invalid_quantity_values"
    mixed = pd.DataFrame({"ItemDescription": ["x", 2], "Specification": ["steel", "steel"]})
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=mixed))
    assert (await service._validate_data_quality(b"data"))["error_type"] == "data_quality_issues"
    assert validation_mod.ExcelValidationService()._validate_data_types(pd.DataFrame({"Unnamed: 0": [1, "bad"]})) == []
    type_issues = validation_mod.ExcelValidationService()._validate_data_types(pd.DataFrame({"Quantity": [1.5, "bad"], "Description": ["x" * 201, "y"], "Uom": ["u" * 21, "u"]}))
    assert any("decimal" in issue for issue in type_issues) and any("Invalid quantity" in issue for issue in type_issues)


# Exit ----------------------------------------------------------------------
class RedisSessionFake:
    def __init__(self, exists=False):
        self.delete_session = AsyncMock(return_value=True)
        self.session_exists = AsyncMock(return_value=exists)


class RedisBaseFake:
    def __init__(self):
        self.delete_pattern = AsyncMock(return_value=1)
        self.delete = AsyncMock()
        self.init_client = AsyncMock()
        self.client = SimpleNamespace(
            scan=AsyncMock(return_value=(0, ["welcome_msg:123456", "other:123456"])),
            delete=AsyncMock(return_value=1),
        )


@pytest.mark.asyncio
async def test_exit_intent_confirmation_cleanup_and_last_message_edges(monkeypatch):
    monkeypatch.setattr(exit_mod, "restore_last_bot_message", AsyncMock())
    service = exit_mod.ExitService.__new__(exit_mod.ExitService)
    service.settings = settings(procucev_link="https://procucev.example")
    service.whatsapp_service = SimpleNamespace(
        send_configurable_buttons=AsyncMock(return_value=MessageResponse(True)),
        send_message=AsyncMock(return_value=MessageResponse(True)),
    )
    service.authentication_service = SimpleNamespace(clear_user_token=AsyncMock(return_value=True))
    service.session_manager = SimpleNamespace(save_session=AsyncMock())
    service.db_manager = SimpleNamespace(append_session_data=MagicMock())

    service._send_exit_confirmation_message = AsyncMock(return_value=True)
    pending = session(conversation_history={"messages": [{"role": "user", "content": "hello"}]}, workflow_state=None)
    result = await service.handle_exit_intent("+1", pending)
    assert result["status"] == "exit_confirmation_pending"
    assert "last_bot_message_before_exit" not in pending.workflow_state
    service.handle_exit_confirmation = AsyncMock(return_value={"status": "confirmed"})
    waiting = session(workflow_state={"exit_pending": True})
    assert (await service.handle_exit_intent("+1", waiting, message={"type": "button_reply", "button_reply": {"id": "confirm_exit", "title": "Yes"}}))["status"] == "confirmed"
    assert service.handle_exit_confirmation.await_args.args[2] is True
    assert (await service.handle_exit_intent("+1", waiting, message="no"))["status"] == "confirmed"
    assert service.handle_exit_confirmation.await_args.args[2] is False
    service.handle_exit_confirmation = AsyncMock(return_value={"status": "direct"})
    assert (await service.handle_exit_intent("+1", pending, show_message=False))["status"] == "direct"
    assert service.handle_exit_confirmation.await_args.args[2] is True

    assert await service._clear_session_data(None)
    redis_session = RedisSessionFake()
    redis_base = RedisBaseFake()
    monkeypatch.setattr("app.redis_db.get_session_redis_service", lambda: redis_session)
    monkeypatch.setattr("app.redis_db.get_redis_service", lambda: redis_base)
    monkeypatch.setattr("app.config.get_settings", lambda: settings(redis_session_storage_enabled=True))
    value = session()
    assert await service._clear_session_data(value)
    assert redis_session.delete_session.await_count == 2
    assert redis_base.client.delete.await_count == 1
    assert redis_base.client.delete.call_args.args[0] == "other:123456"

    monkeypatch.setattr("app.config.get_settings", lambda: settings(redis_session_storage_enabled=False))
    assert await service._clear_session_data(session())
    redis_session.session_exists.side_effect = [True, True]
    monkeypatch.setattr("app.config.get_settings", lambda: settings(redis_session_storage_enabled=True))
    assert not await service._clear_session_data(session())
    redis_session.session_exists.side_effect = None
    redis_base.init_client.side_effect = RuntimeError("scan unavailable")
    assert await service._clear_session_data(session())

    assert await service._send_goodbye_message("+1")
    service.whatsapp_service.send_message.side_effect = RuntimeError("whatsapp")
    assert not await service._send_goodbye_message("+1")
    assert service._get_last_bot_message(session(conversation_history={"messages": [{"role": "user", "content": "x"}]})) is None
    assert service._get_last_bot_message(SimpleNamespace(conversation_history=object())) is None

    service._clear_session_data = AsyncMock(return_value=True)
    service._send_goodbye_message = AsyncMock(return_value=True)
    service.authentication_service.clear_user_token.return_value = True
    service.handle_exit_confirmation = exit_mod.ExitService.handle_exit_confirmation.__get__(service)
    completed = await service.handle_exit_confirmation("+1", session(workflow_state={"exit_pending": True}), True, show_message=False)
    assert completed["status"] == "exit_completed" and not completed["goodbye_sent"]
    declined = await service.handle_exit_confirmation("+1", session(workflow_state={"exit_pending": True}), False)
    assert declined["status"] == "exit_aborted"
    service._send_exit_confirmation_message = exit_mod.ExitService._send_exit_confirmation_message.__get__(service)
    service.whatsapp_service.send_configurable_buttons.side_effect = RuntimeError("buttons")
    assert not await service._send_exit_confirmation_message("+1")


@pytest.mark.asyncio
async def test_excel_validation_additional_content_and_branch_shapes(monkeypatch):
    service = validation_mod.ExcelValidationService()

    assert service._validate_excel_content(b"column_a,column_b\n1,2")["error_type"] == "csv_format_detected"
    assert service._validate_excel_content(b"not-an-excel-file")["error_type"] == "invalid_format"

    blank_row = [SimpleNamespace(value=None), SimpleNamespace(value="")]
    filled_row = [SimpleNamespace(value="Item"), SimpleNamespace(value="Quantity")]
    worksheet = SimpleNamespace(
        iter_rows=lambda: [blank_row, filled_row],
        merged_cells=SimpleNamespace(ranges=[]),
        row_dimensions={},
        column_dimensions={},
        _pivots=[],
    )
    workbook = SimpleNamespace(worksheets=[worksheet], active=worksheet, close=MagicMock())
    monkeypatch.setattr(validation_mod, "load_workbook", Mock(return_value=workbook))
    assert (await service._validate_excel_structure_comprehensive(b"data"))["valid"]

    width = service.MAX_COLUMNS + 1
    wide_frame = pd.DataFrame(
        [[1] * width, ["text"] * width],
        columns=[f"Column {index}" for index in range(width)],
    )
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=wide_frame))
    assert (await service._validate_data_quality(b"data"))["error_type"] == "data_quality_issues"

    unnamed_frame = pd.DataFrame(
        [["Header A", "Header B"], ["pump", "steel"]],
        columns=["Unnamed: 0", "Unnamed: 1"],
    )
    monkeypatch.setattr(validation_mod.pd, "read_excel", Mock(return_value=unnamed_frame))
    assert (await service._validate_data_quality(b"data"))["valid"]

    type_issues = service._validate_data_types(
        pd.DataFrame(
            {
                "Quantity": [1, "2", datetime(2025, 1, 1)],
                "Description": ["pump", "bolt", "nut"],
            }
        )
    )
    assert any("mixed data types" in issue for issue in type_issues)
    assert any("Invalid quantity" in issue for issue in type_issues)


@pytest.mark.asyncio
async def test_exit_alternate_cleanup_and_confirmation_branches(monkeypatch):
    monkeypatch.setattr(exit_mod, "restore_last_bot_message", AsyncMock())
    service = exit_mod.ExitService.__new__(exit_mod.ExitService)
    service.settings = settings(procucev_link="https://procucev.example")
    service.whatsapp_service = SimpleNamespace(
        send_configurable_buttons=AsyncMock(return_value=MessageResponse(True)),
        send_message=AsyncMock(return_value=MessageResponse(True)),
    )
    service.authentication_service = None
    service.session_manager = None
    service.db_manager = SimpleNamespace(append_session_data=MagicMock())
    service._send_exit_confirmation_message = exit_mod.ExitService._send_exit_confirmation_message.__get__(service)

    pending = session(
        workflow_state={},
        conversation_history={"messages": [{"role": "assistant", "content": "last bot message"}]},
    )
    result = await service.handle_exit_intent("+1", pending)
    assert result["status"] == "exit_confirmation_pending"
    assert result["confirmation_sent"] is True
    assert pending.workflow_state["last_bot_message_before_exit"] == "last bot message"

    redis_session = RedisSessionFake(exists=False)
    redis_session.delete_session.side_effect = [False, False]
    redis_base = RedisBaseFake()
    redis_base.client.scan.side_effect = [
        (1, []),
        (0, ["welcome_msg:123456"]),
    ]
    monkeypatch.setattr("app.redis_db.get_session_redis_service", lambda: redis_session)
    monkeypatch.setattr("app.redis_db.get_redis_service", lambda: redis_base)
    monkeypatch.setattr("app.config.get_settings", lambda: settings(redis_session_storage_enabled=True))
    assert await service._clear_session_data(session(workflow_state={"existing": True}))
    assert redis_base.client.delete.await_count == 0

    redis_session.delete_session.side_effect = None
    redis_session.delete_session.return_value = False
    redis_session.session_exists.side_effect = [True, False]
    redis_base.client.scan.side_effect = None
    redis_base.client.scan.return_value = (0, [])
    assert await service._clear_session_data(session())
    redis_base.delete.assert_awaited_once_with("session:s1")

    declined = await service.handle_exit_confirmation(
        "+1", session(workflow_state={"exit_pending": True}), False
    )
    assert declined["status"] == "exit_aborted"

    service._send_exit_confirmation_message = AsyncMock(side_effect=RuntimeError("buttons"))
    failed = await service.handle_exit_intent("+1", session(workflow_state={"other": True}))
    assert failed["status"] == "exit_error"
